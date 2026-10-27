"""Page-aware parsing and chunking — pure functions, no database."""

import re
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

from app.modules.documents.chunking import split_pages_into_chunks
from app.modules.documents.parsers import (
    NO_PAGE,
    PageText,
    pages_to_display_text,
    parse_document_pages,
)


def make_pdf(path: Path, pages: list[str]) -> Path:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=9)
    doc.save(path)
    doc.close()
    return path


def words(text: str) -> list[str]:
    return re.findall(r"\S+", text)


# --- parsing -----------------------------------------------------------------


def test_pdf_pages_are_numbered_from_one(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "a.pdf", ["alpha page", "beta page", "gamma page"])
    pages = parse_document_pages(str(pdf), "pdf")
    assert [p.page for p in pages] == [1, 2, 3]
    assert "alpha" in pages[0].text and "beta" in pages[1].text and "gamma" in pages[2].text


def test_pdf_blank_page_keeps_numbering_aligned(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "a.pdf", ["first", "", "third"])
    pages = parse_document_pages(str(pdf), "pdf")
    assert [p.page for p in pages] == [1, 2, 3]
    assert not pages[1].text.strip()


def test_text_and_markdown_have_no_page_numbers(tmp_path: Path) -> None:
    for ext in ("txt", "md"):
        f = tmp_path / f"a.{ext}"
        f.write_text("hello world", encoding="utf-8")
        assert parse_document_pages(str(f), ext) == [PageText(NO_PAGE, "hello world")]


def test_docx_paragraphs_joined_without_page_numbers(tmp_path: Path) -> None:
    f = tmp_path / "a.docx"
    d = DocxDocument()
    d.add_paragraph("First paragraph")
    d.add_paragraph("")
    d.add_paragraph("Second paragraph")
    d.save(f)
    assert parse_document_pages(str(f), "docx") == [
        PageText(NO_PAGE, "First paragraph\nSecond paragraph")
    ]


def test_display_text_marks_pdf_pages_and_skips_blank_ones() -> None:
    text = pages_to_display_text([PageText(1, "one"), PageText(2, "  "), PageText(3, "three")])
    assert "--- Page 1 ---" in text and "--- Page 3 ---" in text
    assert "--- Page 2 ---" not in text


def test_display_text_for_unpaged_formats_has_no_markers() -> None:
    assert pages_to_display_text([PageText(NO_PAGE, "plain")]) == "plain"


# --- chunking ----------------------------------------------------------------


def test_chunks_never_cross_page_boundaries() -> None:
    chunks = split_pages_into_chunks([PageText(1, "a b c"), PageText(2, "d e"), PageText(3, "f")])
    assert [(c.content, c.page_start, c.page_end) for c in chunks] == [
        ("a b c", 1, 1),
        ("d e", 2, 2),
        ("f", 3, 3),
    ]


def test_blank_pages_do_not_affect_page_ranges() -> None:
    chunks = split_pages_into_chunks([PageText(1, "  "), PageText(2, "only text here"), PageText(3, "")])
    assert [(c.page_start, c.page_end) for c in chunks] == [(2, 2)]


def test_no_text_at_all_produces_no_chunks() -> None:
    assert split_pages_into_chunks([PageText(1, ""), PageText(2, "   \n")]) == []


def test_long_pages_get_their_own_page_numbers() -> None:
    pages = [PageText(n, " ".join(f"p{n}w{i}" for i in range(300))) for n in range(1, 5)]
    chunks = split_pages_into_chunks(pages, chunk_chars=500, overlap_chars=50)

    assert len(chunks) > 4
    for chunk in chunks:
        assert chunk.page_start == chunk.page_end
        # every word in a chunk really comes from the pages the chunk claims
        word_pages = {int(w.split("w")[0][1:]) for w in words(chunk.content)}
        assert word_pages <= set(range(chunk.page_start, chunk.page_end + 1))
    assert chunks[0].page_start == 1 and chunks[-1].page_end == 4
    assert [c.page_start for c in chunks] == sorted(c.page_start for c in chunks)


def test_chunks_cover_every_word_without_cutting_words() -> None:
    source = [PageText(1, " ".join(f"word{i}" for i in range(2000)))]
    original_words = set(words(source[0].text))
    chunks = split_pages_into_chunks(source, chunk_chars=700, overlap_chars=100)

    seen: set[str] = set()
    for chunk in chunks:
        chunk_words = words(chunk.content)
        assert set(chunk_words) <= original_words, "a word was cut in half"
        seen.update(chunk_words)
    assert seen == original_words


def test_consecutive_chunks_overlap() -> None:
    source = [PageText(1, " ".join(f"w{i}" for i in range(1500)))]
    chunks = split_pages_into_chunks(source, chunk_chars=600, overlap_chars=120)
    for a, b in zip(chunks, chunks[1:]):
        assert words(b.content)[0] in words(a.content)
